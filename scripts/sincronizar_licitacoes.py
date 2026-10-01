#!/usr/bin/env python3
"""Leva o que o robô produziu no Pipeline_Licitacoes.xlsx para a planilha Google
"Pipeline Licitações TRX (fonte da verdade)", aba Licitacoes.

Por que existe: as tarefas do Mac (quadro nacional, radar de licitações, monitor de
e-mail) continuam gravando o .xlsx do jeito de sempre. A equipe passa a trabalhar na
planilha Google (direto ou pelo app Central TRX · Edição). Este script roda no GitHub
a cada 10 minutos e junta as duas coisas.

Regras (as mesmas do abastecer_licitacoes.py):
  - A chave é a coluna ID.
  - Colunas do robô (órgão, prazo, valor, links, score...) são atualizadas quando o
    robô traz valor. Valor vazio do robô nunca apaga valor que já existe.
  - Colunas da equipe (Fase, Situacao do envio, Proxima acao, Quando, Responsavel,
    Notas da Carla, Resultado) só são preenchidas quando estão VAZIAS na planilha Google.
    Exceção: "Situacao do envio" ainda em estado do robô é recalculada.
  - Linha nova do robô entra no fim. Linha que sumiu do robô vira Trilha ARQUIVADO
    (nada é apagado).
  - A Trilha (CREDENCIAMENTO, PREGAO, SEM PROCESSO, DESCARTADOS) é recalculada com a
    Fase da planilha Google, então a Carla descartar uma linha tira ela da fila.
  - Só as células que mudaram são gravadas, para não atropelar quem está editando.

Variável de ambiente GOOGLE_SA_JSON: chave da conta de serviço (editora da planilha).
Uso: python3 sincronizar_licitacoes.py Pipeline_Licitacoes.xlsx
"""
import datetime as dt, json, os, re, sys, unicodedata

SHEET_ID = "1ReCnYKNThynTD6-xxvyuDKQ33pakPrZBDqg_1U0PPe8"
ABA = "Licitacoes"
TZ = dt.timezone(dt.timedelta(hours=-3))

HUMANAS = ["Fase", "Situacao do envio", "Proxima acao", "Quando", "Responsavel",
           "Notas da Carla", "Resultado"]
ENVIO_DO_ROBO = {"", "Aguardando envio", "Sem e-mail (achar contato)", "Rascunho gerado"}
ABAS_XLSX = ["CREDENCIAMENTO", "PREGAO", "SEM PROCESSO", "DESCARTADOS"]
CONTROLE = ["Trilha", "ID", "Sincronizado em"]
# Medição da Dor feita direto na planilha Google (tarefa na nuvem ou pesquisa manual).
# Quando existe, ela manda no Score e na Faixa da linha, e o xlsx do robô não mexe neles.
MEDICAO = ["Eixo Janela", "Sinal de compra", "Eixo Necessidade", "Sinal operacional", "Eixo Dor",
           "O que reclamam", "Cobertura", "Medido em", "Fonte da medicao"]
CALCULADAS = ["Score", "Faixa"]
ORIGENS_NUVEM = {"Radar PNCP (GitHub)", "Quadro (nuvem)"}
# Desde 01/10/2026 o quadro público é mantido na nuvem (tarefa "Incremento quadro" e Monitor de e-mail
# gravam direto na planilha). As linhas "quadro-..." do xlsx antigo deixam de sobrescrever a planilha.
QUADRO_NA_NUVEM = True
FASES_INICIAIS = {"", "mapeado", "lendo edital", "montando habilitacao", "em aberto"}


def slug(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")[:60]


def norm(v):
    """Valor comparável e gravável: número fica número, o resto vira texto limpo."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def igual(a, b):
    return str(norm(a)) == str(norm(b))


def prazo_passou(r):
    if str(r.get("Dias") or "").strip().lower() == "encerrado":
        return True
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})", str(r.get("Prazo") or ""))
    if not m:
        return False
    d = dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return d < dt.datetime.now(TZ).date()


def trilha_de(r):
    """Regra do abastecer_licitacoes.py, mais ENCERRADO: processo com prazo vencido em que a
    equipe não avançou (fase inicial) sai da fila de trabalho. Se a Carla já enviou
    documentos ou proposta, a linha continua na trilha do processo."""
    t = _trilha_base(r)
    if t in ("CREDENCIAMENTO", "PREGAO") and prazo_passou(r) and \
            str(r.get("Fase") or "").strip().lower() in FASES_INICIAIS:
        return "ENCERRADO"
    return t


def _num(v):
    try:
        s = str(v).strip().replace(",", ".")
        return float(s) if s != "" else None
    except ValueError:
        return None


def classificar_publico(r, trilha):
    """Regra da Planilha da Dor para instituição PÚBLICA (build_dor.py):
    SCORE = 50% JANELA + 30% NECESSIDADE + 20% DOR, média só dos eixos medidos,
    tetos: sem JANELA 60; só DOR 45; só JANELA 75. Processo aberto = JANELA 100."""
    w, n, d = _num(r.get("Eixo Janela")), _num(r.get("Eixo Necessidade")), _num(r.get("Eixo Dor"))
    if w is None and trilha in ("CREDENCIAMENTO", "PREGAO") and not prazo_passou(r):
        w = 100.0
    pares = [(0.50, w), (0.30, n), (0.20, d)]
    med = [(p, v) for p, v in pares if v is not None]
    if not med:
        return None, "SEM MEDICAO", "0/3"
    sc = sum(p * v for p, v in med) / sum(p for p, _ in med)
    if w is None: sc = min(sc, 60)
    if w is None and n is None: sc = min(sc, 45)
    if w is not None and n is None and d is None: sc = min(sc, 75)
    sc = int(round(sc))
    fx = "ATACAR AGORA" if sc >= 70 else "QUENTE" if sc >= 50 else "MORNO" if sc >= 30 else "FRIO"
    return sc, fx, "%d/3" % len(med)


def medido(r):
    return any(str(r.get(c) or "").strip() for c in ("Eixo Janela", "Eixo Necessidade", "Eixo Dor"))


def _trilha_base(r):
    fase = str(r.get("Fase") or "").strip().lower()
    if fase.startswith(("descartad", "perdid", "sem interesse")):
        return "DESCARTADOS"
    if str(r.get("Origem") or "").startswith("Quadro") and not str(r.get("Edital") or "").strip():
        return "SEM PROCESSO"
    m = slug(r.get("Modalidade"))
    if any(k in m for k in ("credenciament", "chamament", "inexigibil", "manifestacao")):
        return "CREDENCIAMENTO"
    if any(k in m for k in ("pregao", "concorrencia", "dispensa", "leilao", "tomada", "convite", "cotacao")):
        return "PREGAO"
    return "CREDENCIAMENTO" if "lead" in m else "PREGAO"


def ler_xlsx(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for aba in ABAS_XLSX:
        if aba not in wb.sheetnames:
            continue
        rows = list(wb[aba].iter_rows(values_only=True))
        if not rows or not rows[0] or rows[0][0] != "ID":
            continue
        H = [str(h).strip() if h else "" for h in rows[0]]
        for r in rows[1:]:
            if not r or not r[0]:
                continue
            d = {H[i]: r[i] for i in range(min(len(H), len(r))) if H[i]}
            out[str(r[0]).strip()] = d
    return out


def col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def main():
    if len(sys.argv) < 2:
        sys.exit("uso: sincronizar_licitacoes.py Pipeline_Licitacoes.xlsx")
    robo = ler_xlsx(sys.argv[1])
    if len(robo) < 100:
        sys.exit(f"ERRO: o xlsx trouxe só {len(robo)} linhas. Nada foi gravado.")

    from google.oauth2 import service_account
    from googleapiclient.discovery import build
    info = json.loads(os.environ["GOOGLE_SA_JSON"])
    cred = service_account.Credentials.from_service_account_info(
        info, scopes=["https://www.googleapis.com/auth/spreadsheets"])
    api = build("sheets", "v4", credentials=cred, cache_discovery=False).spreadsheets()

    vals = api.values().get(spreadsheetId=SHEET_ID, range=f"{ABA}!A1:AZ",
                            valueRenderOption="UNFORMATTED_VALUE").execute().get("values", [])
    if not vals:
        sys.exit("ERRO: a aba Licitacoes está sem cabeçalho.")
    H = [str(h).strip() for h in vals[0]]
    for c in CONTROLE:
        if c not in H:
            sys.exit(f"ERRO: falta a coluna {c} no cabeçalho da planilha.")
    ix = {h: i for i, h in enumerate(H) if h}
    linhas = vals[1:]
    por_id = {}
    for n, r in enumerate(linhas, start=2):
        r = r + [""] * (len(H) - len(r))
        rid = str(r[ix["ID"]]).strip()
        if rid:
            por_id[rid] = (n, {h: r[i] for h, i in ix.items()})

    agora = dt.datetime.now(TZ).strftime("%d/%m/%Y %H:%M")
    mudancas, novas = [], []
    cont = {"atualizadas": 0, "novas": 0, "arquivadas": 0, "celulas": 0}
    dados_cols = [h for h in H if h not in CONTROLE and h not in MEDICAO]

    for rid, d in robo.items():
        if QUADRO_NA_NUVEM and rid.startswith("quadro-"):
            continue
        if rid in por_id:
            n, cur = por_id[rid]
            merged = dict(cur)
            alt = {}
            ja_medido = medido(cur)
            for c in dados_cols:
                if c not in d:
                    continue
                if ja_medido and c in CALCULADAS:
                    continue
                novo, velho = norm(d.get(c)), norm(cur.get(c))
                if c in HUMANAS:
                    if c == "Situacao do envio" and str(velho) in ENVIO_DO_ROBO:
                        ok = not igual(novo, velho)
                    else:
                        ok = velho == "" and novo != ""
                else:
                    ok = novo != "" and not igual(novo, velho)
                if ok:
                    alt[c] = novo
                    merged[c] = novo
            t = trilha_de(merged)
            if t != str(cur.get("Trilha") or ""):
                alt["Trilha"] = t
            if ja_medido:
                sc, fx, cob = classificar_publico(merged, t)
                for c, v in (("Score", "" if sc is None else sc), ("Faixa", fx), ("Cobertura", cob)):
                    if c in ix and not igual(v, cur.get(c)):
                        alt[c] = v
            if alt:
                alt["Sincronizado em"] = agora
                cont["atualizadas"] += 1
                for c, v in alt.items():
                    mudancas.append({"range": f"{ABA}!{col(ix[c] + 1)}{n}", "values": [[v]]})
        else:
            reg = {c: norm(d.get(c)) for c in dados_cols}
            reg["ID"] = rid
            reg["Trilha"] = trilha_de(reg)
            reg["Sincronizado em"] = agora
            novas.append([reg.get(h, "") for h in H])
            cont["novas"] += 1

    for rid, (n, cur) in por_id.items():
        if rid in robo:
            continue
        if str(cur.get("Origem") or "") in ORIGENS_NUVEM or (QUADRO_NA_NUVEM and rid.startswith("quadro-")):
            # linha criada na nuvem (radar_pncp.py): não depende do xlsx do Mac
            alt = {}
            t = trilha_de(cur)
            if t != str(cur.get("Trilha") or ""):
                alt["Trilha"] = t
            if medido(cur):
                sc, fx, cob = classificar_publico(cur, t)
                for c, v in (("Score", "" if sc is None else sc), ("Faixa", fx), ("Cobertura", cob)):
                    if c in ix and not igual(v, cur.get(c)):
                        alt[c] = v
            if alt:
                alt["Sincronizado em"] = agora
                cont["atualizadas"] += 1
                for c, v in alt.items():
                    mudancas.append({"range": f"{ABA}!{col(ix[c] + 1)}{n}", "values": [[v]]})
            continue
        if str(cur.get("Trilha") or "") != "ARQUIVADO":
            mudancas.append({"range": f"{ABA}!{col(ix['Trilha'] + 1)}{n}", "values": [["ARQUIVADO"]]})
            mudancas.append({"range": f"{ABA}!{col(ix['Sincronizado em'] + 1)}{n}", "values": [[agora]]})
            cont["arquivadas"] += 1

    cont["celulas"] = len(mudancas)
    for i in range(0, len(mudancas), 500):
        api.values().batchUpdate(spreadsheetId=SHEET_ID, body={
            "valueInputOption": "RAW", "data": mudancas[i:i + 500]}).execute()
    if novas:
        api.values().append(spreadsheetId=SHEET_ID, range=f"{ABA}!A1",
                            valueInputOption="RAW", insertDataOption="INSERT_ROWS",
                            body={"values": novas}).execute()
    print("Sincronização de licitações:", json.dumps(cont, ensure_ascii=False),
          f"| robô {len(robo)} linhas | planilha {len(por_id)} antes")


if __name__ == "__main__":
    main()
