#!/usr/bin/env python3
"""Central TRX: gera dados_central.js a partir das planilhas-fonte.

Uso:
  python3 build_central.py --pipeline PIPELINE.json --licitacoes Pipeline_Licitacoes.xlsx \
      --saida dados_central.js [--mod-pipeline ISO] [--mod-licitacoes ISO]

PIPELINE.json: resultado de get_values("Pipeline!A1:AP") do conector Google Sheets
  (objeto com a chave "values", ou a própria lista de linhas).
Pipeline_Licitacoes.xlsx: arquivo baixado do Drive (base64 já decodificado).

A página só mostra dados. Telefone, WhatsApp e e-mail não entram: o contato
continua na planilha, que é a fonte da verdade.
"""
import argparse, datetime as dt, html, json, re, sys, unicodedata

TZ = dt.timezone(dt.timedelta(hours=-3))
PIPELINE_ID = "1rdhgZ_ps8Ih-wwj1dF4WuQhVsz_CnCLJI8JJwTyK9u0"
PIPELINE_GID = 1118277134
LICIT_ID = "1FTHl-0FePl8aSIwGEuzFAj-P1GWEx_H1"
LICIT_SHEET = "1ReCnYKNThynTD6-xxvyuDKQ33pakPrZBDqg_1U0PPe8"
LICIT_GID = 1773203614


def txt(v, lim=None):
    if v is None:
        return ""
    s = str(v)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s)
    s = re.sub(r"\s+", " ", s).strip()
    if lim and len(s) > lim:
        s = s[: lim - 1].rstrip() + "…"
    return s


def num(v):
    if v is None or v == "":
        return None
    s = str(v).strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def data_iso(v):
    """Aceita datetime, 'DD/MM/AAAA', 'DD/MM/AAAA HH:MM' ou 'AAAA-MM-DD'."""
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%dT%H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    m = re.match(r"(\d{2})/(\d{2})/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?", s)
    if m:
        d, mo, y, h, mi = m.groups()
        return f"{y}-{mo}-{d}" + (f"T{int(h):02d}:{mi}" if h else "")
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2}))?", s)
    if m:
        y, mo, d, h, mi = m.groups()
        return f"{y}-{mo}-{d}" + (f"T{h}:{mi}" if h else "")
    return ""


def clinicas(values):
    H = values[0]
    col = {h: i for i, h in enumerate(H) if h}
    def g(r, name, lim=None):
        i = col.get(name)
        return txt(r[i], lim) if i is not None and i < len(r) else ""
    out = []
    for n, r in enumerate(values[1:], start=2):
        if not r or not str(r[0]).strip():
            continue
        out.append({
            "l": n,
            "n": g(r, "Clinica"),
            "c": g(r, "Cidade"),
            "uf": g(r, "UF"),
            "pr": g(r, "Prioridade").upper(),
            "mod": g(r, "Modalidades", 160),
            "g": g(r, "Gancho", 420),
            "resp": g(r, "Responsavel"),
            "fase": g(r, "Fase") or "Sem fase",
            "toque": g(r, "Toque_atual"),
            "prox": data_iso(g(r, "Proximo_toque")),
            "ini": data_iso(g(r, "Inicio_cadencia")),
            "ult": data_iso(g(r, "Ultima_atualizacao")),
            "res": g(r, "Resultado", 200),
            "notas": g(r, "Notas", 420),
            "acao": g(r, "Proxima_Acao", 200),
            "sc": num(g(r, "Score")),
            "fx": g(r, "Faixa") or "SEM MEDICAO",
            "dor": num(g(r, "Dor")),
            "cob": g(r, "Cobertura"),
            "jan": num(g(r, "Janela")),
            "perf": num(g(r, "Perfil")),
            "sjan": g(r, "Sinal_de_compra", 300),
            "nec": num(g(r, "Necessidade")),
            "snec": g(r, "Sinal_operacional", 200),
            "ra": g(r, "Nota_RA", 160),
            "ggl": g(r, "Nota_Google"),
            "oq": g(r, "O_que_reclamam", 360),
            "med": data_iso(g(r, "Medido_em")),
        })
    return out


def abas_licitacoes(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = {}
    for aba in ("CREDENCIAMENTO", "PREGAO", "SEM PROCESSO", "DESCARTADOS"):
        if aba not in wb.sheetnames:
            res[aba] = []
            continue
        rows = list(wb[aba].iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows) if r and sum(1 for c in r if c) > 5), 0)
        H = [str(h).strip() if h else "" for h in rows[hi]]
        res[aba] = [(i + hi + 2, dict(zip(H, r))) for i, r in enumerate(rows[hi + 1:]) if r and any(r)]
    return res


def abas_da_planilha(path):
    """licitacoes.json (get_values da aba Licitacoes) agrupado pela coluna Trilha."""
    pv = json.load(open(path, encoding="utf-8"))
    rows = pv["values"] if isinstance(pv, dict) else pv
    H = [str(h).strip() for h in rows[0]]
    res = {k: [] for k in ("CREDENCIAMENTO", "PREGAO", "SEM PROCESSO", "DESCARTADOS", "ENCERRADO")}
    for n, r in enumerate(rows[1:], start=2):
        d = dict(zip(H, r + [""] * (len(H) - len(r))))
        t = str(d.get("Trilha") or "").strip()
        if t in res and str(d.get("ID") or "").strip():
            res[t].append((n, d))
    return res


def licitacao(linha, d, aba):
    return {
        "l": linha, "aba": aba,
        "id": txt(d.get("ID")),
        "fase": txt(d.get("Fase")) or "Mapeado",
        "enc": txt(d.get("Encaixe")).lower(),
        "sc": num(d.get("Score")),
        "fx": txt(d.get("Faixa")),
        "prazo": data_iso(d.get("Prazo")),
        "org": txt(d.get("Orgao"), 140),
        "c": txt(d.get("Cidade")),
        "uf": txt(d.get("UF")),
        "mod": txt(d.get("Modalidade")),
        "ed": txt(d.get("Edital"), 120),
        "obj": txt(d.get("Objeto"), 420),
        "val": num(d.get("Valor estimado")),
        "onde": txt(d.get("Onde disputa"), 80),
        "bloq": txt(d.get("Bloqueio"), 220),
        "acao": txt(d.get("Proxima acao"), 220),
        "qd": data_iso(d.get("Quando")),
        "resp": txt(d.get("Responsavel")),
        "notas": txt(d.get("Notas da Carla"), 360),
        "res": txt(d.get("Resultado"), 160),
        "led": txt(d.get("Link edital")) if str(d.get("Link edital") or "").startswith("http") else "",
        "lpn": txt(d.get("Link PNCP")) if str(d.get("Link PNCP") or "").startswith("http") else "",
        "pasta": txt(d.get("Pasta do processo")) if str(d.get("Pasta do processo") or "").startswith("http") else "",
        "at": data_iso(d.get("Atualizado em")),
        "jan": num(d.get("Eixo Janela")), "sjan": txt(d.get("Sinal de compra"), 200),
        "nec": num(d.get("Eixo Necessidade")), "snec": txt(d.get("Sinal operacional"), 200),
        "dorx": num(d.get("Eixo Dor")), "oq": txt(d.get("O que reclamam"), 240),
        "cob": txt(d.get("Cobertura")), "med": data_iso(d.get("Medido em")),
    }


def quadro(linha, d):
    return {
        "l": linha,
        "id": txt(d.get("ID")),
        "sit": txt(d.get("Situacao do envio")) or "Sem situação",
        "fase": txt(d.get("Fase")) or "Em aberto",
        "enc": txt(d.get("Encaixe")).lower(),
        "sc": num(d.get("Score")),
        "fx": txt(d.get("Faixa")) or "SEM MEDICAO",
        "org": txt(d.get("Orgao"), 140),
        "c": txt(d.get("Cidade")),
        "uf": txt(d.get("UF")),
        "obj": txt(d.get("Objeto"), 300),
        "bloq": txt(d.get("Bloqueio"), 160),
        "cargo": txt(d.get("Cargo"), 60),
        "acao": txt(d.get("Proxima acao"), 160),
        "resp": txt(d.get("Responsavel")),
        "notas": txt(d.get("Notas da Carla"), 300),
        "at": data_iso(d.get("Atualizado em")),
        "jan": num(d.get("Eixo Janela")), "sjan": txt(d.get("Sinal de compra"), 200),
        "nec": num(d.get("Eixo Necessidade")), "snec": txt(d.get("Sinal operacional"), 200),
        "dorx": num(d.get("Eixo Dor")), "oq": txt(d.get("O que reclamam"), 240),
        "cob": txt(d.get("Cobertura")), "med": data_iso(d.get("Medido em")),
    }


DOCS_SHEET = "13nXaAsgQ4czDGnp6SZd2ConSeV_8_1exZG_lH4l9hD8"
DOCS_GID = 423972086


def documentos_alerta(path):
    """Documentos TRX: so o que efetivamente VENCEU (Status VENCIDO) ou esta DESATUALIZADO
    (coluna T preenchida). Vence em 15 ou 30 dias nao entra: o aviso e so para o que trava."""
    try:
        dv = json.load(open(path, encoding="utf-8"))
    except Exception:
        return []
    rows = dv["values"] if isinstance(dv, dict) else dv
    if not rows:
        return []
    H = [str(h).strip() for h in rows[0]]
    out = []
    for n, r in enumerate(rows[1:], start=2):
        d = dict(zip(H, r + [""] * (len(H) - len(r))))
        doc = txt(d.get("Documento"), 140)
        if not doc:
            continue
        venc = str(d.get("Status") or "").strip().upper() == "VENCIDO"
        desat = txt(d.get("Desatualizado (motivo)"), 220)
        if not venc and not desat:
            continue
        out.append({"l": n, "doc": doc, "venc": venc, "val": data_iso(d.get("Validade")),
                    "mot": desat, "resp": txt(d.get("Responsavel")), "sit": txt(d.get("Situacao"), 60)})
    out.sort(key=lambda x: (0 if x["venc"] else 1, x["val"] or "9999"))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pipeline", required=True)
    ap.add_argument("--licitacoes", required=True, help="xlsx do robô (reserva) ou licitacoes.json da planilha Google")
    ap.add_argument("--saida", required=True)
    ap.add_argument("--mod-pipeline", default="")
    ap.add_argument("--mod-licitacoes", default="")
    ap.add_argument("--documentos", default="", help="documentos.json (aba Documentos da planilha Documentos TRX)")
    a = ap.parse_args()

    pv = json.load(open(a.pipeline, encoding="utf-8"))
    values = pv["values"] if isinstance(pv, dict) else pv
    cl = clinicas(values)
    da_planilha = a.licitacoes.endswith(".json")
    ab = abas_da_planilha(a.licitacoes) if da_planilha else abas_licitacoes(a.licitacoes)
    lic = [licitacao(l, d, "CREDENCIAMENTO") for l, d in ab["CREDENCIAMENTO"]] + \
          [licitacao(l, d, "PREGAO") for l, d in ab["PREGAO"]]
    qd = [quadro(l, d) for l, d in ab["SEM PROCESSO"]]

    if len(cl) < 100 or len(qd) < 100:
        sys.exit(f"ERRO: leitura suspeita (clinicas={len(cl)}, quadro={len(qd)}). Nada foi gerado.")

    agora = dt.datetime.now(TZ)
    # processos que encerraram nos últimos 10 dias sem Resultado: a equipe registra o que houve
    lim = (agora - dt.timedelta(days=10)).strftime("%Y-%m-%d")
    enc_rec = [licitacao(l, d, "ENCERRADO") for l, d in ab.get("ENCERRADO", [])]
    enc_rec = sorted([e for e in enc_rec if e["prazo"] and e["prazo"][:10] >= lim and not e["res"]],
                     key=lambda e: e["prazo"], reverse=True)
    # tudo que tem Resultado registrado (licitações de qualquer trilha e clínicas)
    resultados = []
    pv2 = json.load(open(a.licitacoes, encoding="utf-8")) if da_planilha else None
    if pv2:
        rows = pv2["values"] if isinstance(pv2, dict) else pv2
        H2 = [str(h).strip() for h in rows[0]]
        for n, r in enumerate(rows[1:], start=2):
            d = dict(zip(H2, r + [""] * (len(H2) - len(r))))
            res = txt(d.get("Resultado"), 220)
            if res and str(d.get("ID") or "").strip() and not res.startswith("Ruído do radar"):
                resultados.append({"src": "quadro" if str(d.get("Trilha")) == "SEM PROCESSO" else "licitacoes",
                                   "l": n, "n": txt(d.get("Orgao"), 140) or txt(d.get("ID")), "uf": txt(d.get("UF")),
                                   "tp": txt(d.get("Modalidade")) or txt(d.get("Trilha")), "fase": txt(d.get("Fase")),
                                   "res": res, "dt": data_iso(d.get("Prazo")), "resp": txt(d.get("Responsavel"))})
    for c in cl:
        if c.get("res"):
            resultados.append({"src": "clinicas", "l": c["l"], "n": c["n"], "uf": c["uf"], "tp": "Clínica",
                               "fase": c["fase"], "res": c["res"], "dt": c.get("ult", ""), "resp": c["resp"]})
    dados = {
        "gerado": agora.strftime("%Y-%m-%dT%H:%M"),
        "fontes": {
            "pipeline": {"nome": "Pipeline TRX — Prospecção (fonte da verdade)", "id": PIPELINE_ID,
                          "gid": PIPELINE_GID, "mod": a.mod_pipeline, "linhas": len(cl)},
            "licitacoes": {"nome": "Pipeline Licitações TRX (fonte da verdade)" if da_planilha else "Pipeline_Licitacoes.xlsx",
                            "id": LICIT_SHEET if da_planilha else LICIT_ID, "gid": LICIT_GID if da_planilha else None,
                            "sheet": da_planilha, "mod": a.mod_licitacoes,
                            "linhas": len(lic), "quadro": len(qd), "descartados": len(ab["DESCARTADOS"]),
                            "encerrados": len(ab.get("ENCERRADO", []))},
        },
        "clinicas": cl,
        "licitacoes": lic,
        "quadro": qd,
        "encerrados_recentes": enc_rec,
        "resultados": resultados,
        "documentos_alerta": documentos_alerta(a.documentos) if a.documentos else [],
        "documentos_fonte": {"id": DOCS_SHEET, "gid": DOCS_GID},
    }
    js = "window.CENTRAL=" + json.dumps(dados, ensure_ascii=False, separators=(",", ":")) + ";\n"
    open(a.saida, "w", encoding="utf-8").write(js)
    print(f"OK {a.saida}: {len(cl)} clinicas, {len(lic)} licitacoes, {len(qd)} municipios do quadro, "
          f"{len(js)//1024} KB, gerado {dados['gerado']}")


if __name__ == "__main__":
    main()
